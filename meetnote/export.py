"""Export a note as text, Markdown, Word, subtitles or JSON."""
from __future__ import annotations

import io
import json
from datetime import datetime

from .summarize import NOTE_TYPES, fmt_time


def _ctx(note: dict, segments: list[dict], speakers: list[dict]):
    names = {s["idx"]: s["name"] for s in speakers}
    when = datetime.fromtimestamp(note.get("recorded_at") or note["created_at"])
    return names, when


def _date(when: datetime) -> str:
    wd = "월화수목금토일"[when.weekday()]
    return when.strftime(f"%Y.%m.%d ({wd}) %H:%M")


def _dur(sec: float) -> str:
    sec = int(sec or 0)
    h, m, s = sec // 3600, (sec % 3600) // 60, sec % 60
    return f"{h}시간 {m}분 {s}초" if h else f"{m}분 {s}초"


def _used_speakers(segments, names) -> list[str]:
    seen = []
    for s in segments:
        n = names.get(s["speaker"], f"참석자 {s['speaker'] + 1}")
        if n not in seen:
            seen.append(n)
    return seen


def to_txt(note, segments, speakers, include_summary: bool = True) -> str:
    names, when = _ctx(note, segments, speakers)
    lines = [note["title"], f"{_date(when)} · {_dur(note['duration'])}",
             "참석자: " + ", ".join(_used_speakers(segments, names)), ""]
    summ = note.get("summary")
    if include_summary and summ:
        lines += ["[AI 요약]", summ.get("overview", ""), ""]
        if summ.get("decisions"):
            lines += ["[결정 사항]"] + [f"- {d}" for d in summ["decisions"]] + [""]
        if summ.get("action_items"):
            lines += ["[할 일]"] + [
                f"- {a['task']}" + (f" (담당: {a['owner']})" if a.get("owner") else "")
                + (f" (기한: {a['due']})" if a.get("due") else "") for a in summ["action_items"]] + [""]
        lines += ["[전체 대화]", ""]
    for s in segments:
        lines.append(f"{names.get(s['speaker'], '참석자')} {fmt_time(s['start'])}")
        lines.append(s["text"])
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def minutes_md(note, segments, speakers) -> str:
    names, when = _ctx(note, segments, speakers)
    summ = note.get("summary") or {}
    kind = NOTE_TYPES.get(note.get("note_type") or "meeting", "회의")
    out = [f"# {note['title']}", "",
           "| 항목 | 내용 |", "|---|---|",
           f"| 일시 | {_date(when)} |",
           f"| 소요 시간 | {_dur(note['duration'])} |",
           f"| 참석자 | {', '.join(_used_speakers(segments, names))} |",
           f"| 유형 | {kind} |"]
    if summ.get("keywords"):
        out.append(f"| 키워드 | {', '.join(summ['keywords'])} |")
    out.append("")
    if summ.get("one_line"):
        out += ["## 한 줄 요약", "", f"> {summ['one_line']}", ""]
    if summ.get("overview"):
        out += ["## 요약", "", summ["overview"], ""]
    if summ.get("key_points"):
        out += ["## 핵심 내용", ""] + [f"- {p}" for p in summ["key_points"]] + [""]
    if summ.get("topics"):
        out += ["## 주제별 논의 내용", ""]
        for i, t in enumerate(summ["topics"], 1):
            out.append(f"### {i}. {t['title']}" + (f" ({t['time']})" if t.get("time") else ""))
            out += [f"- {p}" for p in t.get("points", [])] + [""]
    if summ.get("decisions"):
        out += ["## 결정 사항", ""] + [f"{i}. {d}" for i, d in enumerate(summ["decisions"], 1)] + [""]
    if summ.get("action_items"):
        out += ["## 할 일 (Action Items)", "", "| 완료 | 할 일 | 담당 | 기한 |", "|---|---|---|---|"]
        for a in summ["action_items"]:
            out.append(f"| {'☑' if a.get('done') else '☐'} | {a['task']} | {a.get('owner') or '-'} | {a.get('due') or '-'} |")
        out.append("")
    if summ.get("open_issues"):
        out += ["## 미결 사항 · 추가 확인 필요", ""] + [f"- {x}" for x in summ["open_issues"]] + [""]
    if summ.get("participants"):
        out += ["## 참석자별 주요 발언", ""] + [f"- **{p['name']}**: {p['summary']}" for p in summ["participants"]] + [""]
    if note.get("memo"):
        out += ["## 메모", "", note["memo"], ""]
    return "\n".join(out)


def to_md(note, segments, speakers) -> str:
    names, _ = _ctx(note, segments, speakers)
    out = [minutes_md(note, segments, speakers), "---", "", "## 전체 대화", ""]
    for s in segments:
        out.append(f"**{names.get(s['speaker'], '참석자')}** `{fmt_time(s['start'])}`  ")
        out.append(s["text"])
        out.append("")
    return "\n".join(out)


def _srt_time(t: float, sep: str = ",") -> str:
    ms = int(round(t * 1000))
    h, ms = divmod(ms, 3_600_000)
    m, ms = divmod(ms, 60_000)
    s, ms = divmod(ms, 1000)
    return f"{h:02d}:{m:02d}:{s:02d}{sep}{ms:03d}"


def _cues(segments, names, max_chars: int = 42):
    """Subtitle cues: split long paragraphs using word times when available."""
    cues = []
    for s in segments:
        who = names.get(s["speaker"], "")
        words = s.get("words") or []
        if not words:
            cues.append((s["start"], s["end"], who, s["text"]))
            continue
        cur = []
        for w in words:
            cur.append(w)
            text = " ".join(x["w"] for x in cur)
            if len(text) >= max_chars or w["w"].endswith((".", "?", "!")):
                cues.append((cur[0]["s"], cur[-1]["e"], who, text))
                cur = []
        if cur:
            cues.append((cur[0]["s"], cur[-1]["e"], who, " ".join(x["w"] for x in cur)))
    return cues


def to_srt(note, segments, speakers, with_speaker: bool = True) -> str:
    names, _ = _ctx(note, segments, speakers)
    out = []
    for i, (a, b, who, text) in enumerate(_cues(segments, names), 1):
        out += [str(i), f"{_srt_time(a)} --> {_srt_time(max(b, a + 0.5))}",
                f"[{who}] {text}" if with_speaker and who else text, ""]
    return "\n".join(out)


def to_vtt(note, segments, speakers) -> str:
    names, _ = _ctx(note, segments, speakers)
    out = ["WEBVTT", ""]
    for a, b, who, text in _cues(segments, names):
        out += [f"{_srt_time(a, '.')} --> {_srt_time(max(b, a + 0.5), '.')}", f"<v {who}>{text}" if who else text, ""]
    return "\n".join(out)


def to_json(note, segments, speakers, bookmarks) -> str:
    keep = {k: note[k] for k in ("id", "title", "created_at", "recorded_at", "duration", "language",
                                 "note_type", "memo", "summary")}
    return json.dumps({"note": keep, "speakers": speakers, "segments": segments, "bookmarks": bookmarks},
                      ensure_ascii=False, indent=2)


def to_docx(note, segments, speakers, include_transcript: bool = True) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    names, when = _ctx(note, segments, speakers)
    summ = note.get("summary") or {}
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Apple SD Gothic Neo"
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), "Apple SD Gothic Neo")

    def shade(cell, color="EEF1F6"):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color)
        tcPr.append(shd)

    doc.add_heading(note["title"], level=0)
    info = [("일시", _date(when)), ("소요 시간", _dur(note["duration"])),
            ("참석자", ", ".join(_used_speakers(segments, names))),
            ("유형", NOTE_TYPES.get(note.get("note_type") or "meeting", "회의"))]
    if summ.get("keywords"):
        info.append(("키워드", ", ".join(summ["keywords"])))
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in info:
        row = t.add_row().cells
        row[0].text, row[1].text = k, v
        row[0].paragraphs[0].runs[0].bold = True
        shade(row[0])
        row[0].width = Pt(80)

    def heading(text):
        doc.add_heading(text, level=1)

    def bullets(items, style_name="List Bullet"):
        for x in items:
            doc.add_paragraph(x, style=style_name)

    if summ.get("one_line"):
        heading("한 줄 요약")
        p = doc.add_paragraph(summ["one_line"])
        p.runs[0].italic = True
    if summ.get("overview"):
        heading("요약")
        doc.add_paragraph(summ["overview"])
    if summ.get("key_points"):
        heading("핵심 내용")
        bullets(summ["key_points"])
    if summ.get("topics"):
        heading("주제별 논의 내용")
        for i, tp in enumerate(summ["topics"], 1):
            doc.add_heading(f"{i}. {tp['title']}" + (f"  ({tp['time']})" if tp.get("time") else ""), level=2)
            bullets(tp.get("points", []))
    if summ.get("decisions"):
        heading("결정 사항")
        bullets(summ["decisions"], "List Number")
    if summ.get("action_items"):
        heading("할 일 (Action Items)")
        at = doc.add_table(rows=1, cols=4)
        at.style = "Table Grid"
        for c, h in zip(at.rows[0].cells, ("완료", "할 일", "담당", "기한")):
            c.text = h
            c.paragraphs[0].runs[0].bold = True
            shade(c)
        for a in summ["action_items"]:
            r = at.add_row().cells
            r[0].text = "☑" if a.get("done") else "☐"
            r[1].text = a["task"]
            r[2].text = a.get("owner") or "-"
            r[3].text = a.get("due") or "-"
    if summ.get("open_issues"):
        heading("미결 사항 · 추가 확인 필요")
        bullets(summ["open_issues"])
    if summ.get("participants"):
        heading("참석자별 주요 발언")
        for p in summ["participants"]:
            para = doc.add_paragraph(style="List Bullet")
            para.add_run(p["name"] + ": ").bold = True
            para.add_run(p["summary"])
    if note.get("memo"):
        heading("메모")
        doc.add_paragraph(note["memo"])
    if include_transcript and segments:
        doc.add_page_break()
        heading("전체 대화")
        for s in segments:
            p = doc.add_paragraph()
            r = p.add_run(names.get(s["speaker"], "참석자"))
            r.bold = True
            r2 = p.add_run("  " + fmt_time(s["start"]))
            r2.font.color.rgb = RGBColor(0x88, 0x8E, 0x99)
            r2.font.size = Pt(9)
            p.paragraph_format.space_after = Pt(0)
            doc.add_paragraph(s["text"]).paragraph_format.space_after = Pt(8)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
